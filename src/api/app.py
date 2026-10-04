"""
API کامل برای داشبورد ANPR - نسخه ۳.۰
"""

from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import HTMLResponse, FileResponse, StreamingResponse
from pydantic import BaseModel
from typing import Optional, List
import os
import io
import csv
import cv2
import time
import json
from datetime import datetime
from persiantools.jdatetime import JalaliDate

from src.database import crud
from src.database.models import backup_db, get_backups, get_db_size
from src.service.cleanup import cleanup_old_snapshots, get_snapshots_size, get_snapshots_count


app = FastAPI(title="ANPR API", version="3.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)


# ==================== تنظیمات دوربین ====================
CAMERAS_FILE = "data/cameras.json"

DEFAULT_CAMERAS = [
    {"id": "cam-1", "name": "دوربین ورودی", "rtsp": "", "enabled": True},
    {"id": "cam-2", "name": "دوربین خروجی", "rtsp": "", "enabled": False},
    {"id": "cam-3", "name": "دوربین ۳", "rtsp": "", "enabled": False},
    {"id": "cam-4", "name": "دوربین ۴", "rtsp": "", "enabled": False},
]


def load_cameras():
    if not os.path.exists(CAMERAS_FILE):
        os.makedirs(os.path.dirname(CAMERAS_FILE), exist_ok=True)
        with open(CAMERAS_FILE, "w", encoding="utf-8") as f:
            json.dump(DEFAULT_CAMERAS, f, ensure_ascii=False, indent=2)
        return DEFAULT_CAMERAS
    with open(CAMERAS_FILE, "r", encoding="utf-8") as f:
        return json.load(f)


def save_cameras(cameras):
    with open(CAMERAS_FILE, "w", encoding="utf-8") as f:
        json.dump(cameras, f, ensure_ascii=False, indent=2)


# ==================== Pydantic ====================
class WhitelistItem(BaseModel):
    plate: str
    plate_fa: str = ""
    owner_name: str = ""
    vehicle_type: str = ""
    department: str = ""
    phone: str = ""


class BlacklistItem(BaseModel):
    plate: str
    plate_fa: str = ""
    reason: str = ""


class CameraItem(BaseModel):
    id: str
    name: str
    rtsp: str = ""
    enabled: bool = True


# ==================== Helpers ====================
def clean_for_json(data):
    if isinstance(data, dict):
        return {k: clean_for_json(v) for k, v in data.items()}
    elif isinstance(data, list):
        return [clean_for_json(item) for item in data]
    elif isinstance(data, bytes):
        try:
            return data.decode("utf-8")
        except UnicodeDecodeError:
            return data.decode("utf-8", errors="replace")
    return data


def safe_float(value, default=0.0):
    if value is None:
        return default
    try:
        return float(value)
    except (ValueError, TypeError):
        return default


# ==================== Dashboard ====================
@app.get("/", response_class=HTMLResponse)
async def dashboard():
    html_path = os.path.join(os.path.dirname(__file__), "dashboard.html")
    if os.path.exists(html_path):
        with open(html_path, "r", encoding="utf-8") as f:
            return f.read()
    return "<h1>داشبورد پیدا نشد</h1>"


# ==================== تصاویر ====================
@app.get("/api/snapshot/{path:path}")
async def get_snapshot(path: str):
    clean_path = path.replace("/", os.sep)
    if not clean_path.startswith("data"):
        clean_path = os.path.join("data", "snapshots", clean_path)
    full_path = os.path.join(os.getcwd(), clean_path)
    if not os.path.exists(full_path):
        raise HTTPException(404, "تصویر پیدا نشد")
    return FileResponse(full_path, media_type="image/jpeg")


# ==================== Stream ====================
def generate_frames(rtsp_url):
    cap = cv2.VideoCapture(rtsp_url)
    if not cap.isOpened():
        return
    try:
        while True:
            ret, frame = cap.read()
            if not ret:
                break
            frame = cv2.resize(frame, (640, 360))
            _, buffer = cv2.imencode('.jpg', frame, [cv2.IMWRITE_JPEG_QUALITY, 70])
            yield (b'--frame\r\n'
                   b'Content-Type: image/jpeg\r\n\r\n' + buffer.tobytes() + b'\r\n')
            time.sleep(0.05)
    finally:
        cap.release()


@app.get("/api/camera/{camera_id}/stream")
async def camera_stream(camera_id: str):
    cameras = load_cameras()
    camera = next((c for c in cameras if c["id"] == camera_id), None)
    if not camera:
        raise HTTPException(404, "دوربین پیدا نشد")
    if not camera.get("rtsp"):
        raise HTTPException(400, "آدرس RTSP تنظیم نشده")
    return StreamingResponse(
        generate_frames(camera["rtsp"]),
        media_type="multipart/x-mixed-replace; boundary=frame"
    )


# ==================== دوربین‌ها ====================
@app.get("/api/cameras")
async def get_cameras():
    return {"count": len(load_cameras()), "cameras": load_cameras()}


@app.post("/api/cameras")
async def add_camera(camera: CameraItem):
    cameras = load_cameras()
    existing = next((c for c in cameras if c["id"] == camera.id), None)
    if existing:
        existing.update(camera.dict())
    else:
        cameras.append(camera.dict())
    save_cameras(cameras)
    return {"status": "ok"}


@app.delete("/api/cameras/{camera_id}")
async def delete_camera(camera_id: str):
    cameras = [c for c in load_cameras() if c["id"] != camera_id]
    save_cameras(cameras)
    return {"status": "ok"}


# ==================== Stats ====================
@app.get("/api/stats")
async def get_stats():
    summary = crud.get_stats_summary()
    today = JalaliDate(datetime.now()).strftime("%Y/%m/%d")
    events = crud.get_events(limit=10000)
    today_events = [e for e in events if e.get("date") == today]

    return {
        **summary,
        "today_events": len(today_events),
        "date": today,
        "time": datetime.now().strftime("%H:%M:%S"),
    }


@app.get("/api/stats/hourly")
async def get_hourly_stats(date: Optional[str] = None):
    return {"date": date or crud.get_now()["date"], "hourly": crud.get_hourly_stats(date)}


@app.get("/api/stats/daily")
async def get_daily_stats(days: int = 7):
    return {"days": days, "daily": crud.get_daily_stats(days)}


@app.get("/api/stats/top-plates")
async def get_top_plates(limit: int = 10):
    return {"top_plates": crud.get_top_plates(limit)}


# ==================== Events ====================
@app.get("/api/events")
async def get_events(
    plate: Optional[str] = None,
    date: Optional[str] = None,
    direction: Optional[str] = None,
    limit: int = 100,
    offset: int = 0,
):
    events = crud.get_events(plate=plate, date=date, direction=direction, limit=limit, offset=offset)
    total = crud.count_events(plate=plate, date=date, direction=direction)
    events = clean_for_json(events)
    return {"count": len(events), "total": total, "events": events}


@app.delete("/api/events/{event_id}")
async def delete_event(event_id: int):
    crud.delete_event(event_id)
    return {"status": "ok"}


# ==================== CSV ====================
@app.get("/api/events/export/csv")
async def export_csv(plate: Optional[str] = None, date: Optional[str] = None):
    events = crud.get_events(plate=plate, date=date, limit=100000)

    output = io.StringIO()
    output.write('\ufeff')
    writer = csv.writer(output)
    writer.writerow(["شناسه", "تاریخ", "ساعت", "پلاک", "مالک", "نوع",
                     "جهت", "مدت", "اطمینان", "دوربین", "تصویر"])

    for ev in events:
        direction_fa = "ورود" if ev.get("direction") == "entry" else "خروج"
        duration = ev.get("duration_min")
        writer.writerow([
            ev.get("id", ""),
            ev.get("date", ""),
            ev.get("time", ""),
            ev.get("plate", ""),
            ev.get("owner_name", "متفرقه"),
            ev.get("vehicle_type", ""),
            direction_fa,
            f"{safe_float(duration):.1f}" if duration else "",
            f"{safe_float(ev.get('confidence')):.2f}",
            ev.get("camera_id", ""),
            ev.get("snapshot_entry") or ev.get("snapshot_exit") or "",
        ])

    filename = f"events_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv; charset=utf-8-sig",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


# ==================== Excel ====================
@app.get("/api/events/export/excel")
async def export_excel(plate: Optional[str] = None, date: Optional[str] = None):
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill

    events = crud.get_events(plate=plate, date=date, limit=100000)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "رویدادها"
    ws.sheet_view.rightToLeft = True

    headers = ["شناسه", "تاریخ", "ساعت", "پلاک", "مالک", "نوع", "جهت", "مدت", "اطمینان", "دوربین"]
    ws.append(headers)

    fill = PatternFill(start_color="3b82f6", end_color="3b82f6", fill_type="solid")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")

    for ev in events:
        direction_fa = "ورود" if ev.get("direction") == "entry" else "خروج"
        duration = ev.get("duration_min")
        ws.append([
            ev.get("id", ""),
            ev.get("date", ""),
            ev.get("time", ""),
            ev.get("plate", ""),
            ev.get("owner_name", "متفرقه"),
            ev.get("vehicle_type", ""),
            direction_fa,
            f"{safe_float(duration):.1f}" if duration else "",
            f"{safe_float(ev.get('confidence')):.2f}",
            ev.get("camera_id", ""),
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max_len + 4

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"events_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


# ==================== Vehicles ====================
@app.get("/api/vehicles")
async def get_vehicles(status: Optional[str] = None):
    vehicles = clean_for_json(crud.get_all_vehicles(status=status))
    return {"count": len(vehicles), "vehicles": vehicles}


# ==================== Whitelist ====================
@app.get("/api/whitelist")
async def get_whitelist():
    items = clean_for_json(crud.get_whitelist())
    return {"count": len(items), "items": items}


@app.post("/api/whitelist")
async def add_whitelist(item: WhitelistItem):
    crud.add_to_whitelist(item.plate, item.plate_fa, item.owner_name,
                          item.vehicle_type, item.department, item.phone)
    return {"status": "ok"}


@app.put("/api/whitelist/{plate}")
async def update_whitelist(plate: str, item: WhitelistItem):
    crud.remove_from_whitelist(plate)
    crud.add_to_whitelist(item.plate, item.plate_fa, item.owner_name,
                          item.vehicle_type, item.department, item.phone)
    return {"status": "ok"}


@app.delete("/api/whitelist/{plate}")
async def remove_whitelist(plate: str):
    crud.remove_from_whitelist(plate)
    return {"status": "ok"}


# ==================== Blacklist ====================
@app.get("/api/blacklist")
async def get_blacklist():
    items = clean_for_json(crud.get_blacklist())
    return {"count": len(items), "items": items}


@app.post("/api/blacklist")
async def add_blacklist(item: BlacklistItem):
    crud.add_to_blacklist(item.plate, item.plate_fa, item.reason)
    return {"status": "ok"}


@app.delete("/api/blacklist/{plate}")
async def remove_blacklist(plate: str):
    crud.remove_from_blacklist(plate)
    return {"status": "ok"}


# ==================== Plate Status ====================
@app.get("/api/plate-status/{plate}")
async def get_plate_status(plate: str):
    """وضعیت پلاک: whitelist / blacklist / unknown"""
    wl = crud.is_in_whitelist(plate)
    if wl:
        return {"status": "whitelist", "info": clean_for_json(wl)}

    bl = crud.is_in_blacklist(plate)
    if bl:
        return {"status": "blacklist", "info": clean_for_json(bl)}

    return {"status": "unknown", "info": None}


# ==================== Backup ====================
@app.get("/api/backup/list")
async def list_backups():
    backups = get_backups()
    return {"count": len(backups), "backups": backups, "db_size": get_db_size()}


@app.post("/api/backup/create")
async def create_backup():
    path = backup_db()
    if path:
        return {"status": "ok", "path": path}
    raise HTTPException(500, "بکاپ ناموفق")


@app.post("/api/backup/restore/{filename}")
async def restore_backup(filename: str):
    from src.database.models import restore_db
    backup_path = os.path.join("data/backups", filename)
    if restore_db(backup_path):
        return {"status": "ok"}
    raise HTTPException(404, "بکاپ پیدا نشد")


# ==================== Cleanup ====================
@app.get("/api/cleanup/size")
async def cleanup_size():
    return {
        "snapshots_size_mb": round(get_snapshots_size() / (1024*1024), 1),
        "snapshots_count": get_snapshots_count(),
    }


@app.post("/api/cleanup/old-snapshots")
async def cleanup_snapshots(days: int = 7):
    deleted = cleanup_old_snapshots(days)
    return {"status": "ok", "deleted": deleted}


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)